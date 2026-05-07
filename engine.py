"""
Motor de agendamento da Sala Transforma.
Carrega dados do Excel + overrides.json e roda o agendamento completo.
"""
import sys
sys.stdout.reconfigure(encoding='utf-8') if hasattr(sys.stdout, 'reconfigure') else None

import os
import re
import json
import pandas as pd
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import PyPDF2

# ============== CONSTANTES ==============
DAYS = ['27/05 (Qua)', '28/05 (Qui)', '29/05 (Sex)']
SLOTS = ['10h-10h40', '11h-11h40', '11h50-12h30', '14h-14h40', '15h-15h40', '16h-16h40', '17h-17h40']
TABLES_PER_SLOT = 13  # capacidade padrão (13 mesas em todos os slots/dias)
TABLES_PER_SLOT_OVERRIDE = {}

def cap_for(day, slot):
    return TABLES_PER_SLOT_OVERRIDE.get((day, slot), TABLES_PER_SLOT)
DAY_COL_RANGES = {
    '27/05 (Qua)': list(range(8, 15)),
    '28/05 (Qui)': list(range(15, 22)),
    '29/05 (Sex)': list(range(22, 29)),
}
BLOCKED_STATUSES = {'rodada', 'buffer', 'indisp', 'pitching', 'viagem', 'painel'}

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_PATH = os.path.join(BASE_DIR, 'Rio2C_Agenda_Matchmaking_SalaTransforma_2026.xlsx')
MATCHES_XLSX_PATH = os.path.join(BASE_DIR, 'matches_aprovados.xlsx')
PDF_PATH = os.path.join(BASE_DIR, 'matches_aprovados_detalhado_20260504_1055.pdf')
OVERRIDES_PATH = os.path.join(BASE_DIR, 'overrides.json')


def find_matches_xlsx():
    """Procura por qualquer arquivo de matches xlsx (com timestamp ou não)."""
    if os.path.exists(MATCHES_XLSX_PATH):
        return MATCHES_XLSX_PATH
    # Procura por padrões matches_aprovados*.xlsx
    import glob
    candidates = sorted(glob.glob(os.path.join(BASE_DIR, 'matches_aprovados*.xlsx')))
    # Filtra "_detalhado" que é versão antiga
    candidates = [c for c in candidates if '_detalhado' not in os.path.basename(c).lower()]
    return candidates[-1] if candidates else None


# ============== GITHUB AUTO-COMMIT ==============
def commit_to_github(file_path, repo_filename, commit_message, token):
    """Faz upload de um arquivo para o repo do GitHub via API.
    Retorna (success: bool, message: str).
    """
    import base64, json, urllib.request, urllib.error
    REPO = 'cainaking/rio2c-agenda'
    try:
        with open(file_path, 'rb') as f:
            content_b64 = base64.b64encode(f.read()).decode()
        # Pega SHA atual se existir
        api_url = f'https://api.github.com/repos/{REPO}/contents/{repo_filename}'
        sha = None
        try:
            req = urllib.request.Request(api_url,
                headers={'Authorization': f'token {token}', 'Accept': 'application/vnd.github+json'})
            with urllib.request.urlopen(req) as r:
                sha = json.loads(r.read()).get('sha')
        except urllib.error.HTTPError as e:
            if e.code != 404: raise
        # PUT
        body = {'message': commit_message, 'content': content_b64}
        if sha: body['sha'] = sha
        req = urllib.request.Request(api_url, method='PUT',
            headers={'Authorization': f'token {token}', 'Accept': 'application/vnd.github+json',
                    'Content-Type': 'application/json'},
            data=json.dumps(body).encode())
        with urllib.request.urlopen(req) as r:
            result = json.loads(r.read())
            sha_short = result['commit']['sha'][:7]
            return True, f"Commitado no GitHub (commit {sha_short})"
    except Exception as e:
        return False, f"Erro ao commitar: {e}"


# ============== OVERRIDES ==============
def load_overrides():
    if os.path.exists(OVERRIDES_PATH):
        with open(OVERRIDES_PATH, 'r', encoding='utf-8') as f:
            return json.load(f)
    return {
        'bloqueios_horario': [],     # [{empresa, dia, slot, motivo}]
        'matches_extras': [],         # [{empresa_a, empresa_b, reps_a, reps_b, tier, score}]
        'matches_cancelados': [],     # [{empresa_a, empresa_b}]
        'correcoes_empresa': [        # [{nome, empresa_antiga, empresa_nova}]
            {'nome': 'tiago camapny', 'empresa_antiga': 'Globoplay', 'empresa_nova': 'Globo Internacional'},
            {'nome': 'andre saad', 'empresa_antiga': 'NewCo', 'empresa_nova': 'Grupo Bandeirantes'},
            {'nome': 'gabriel ferraz', 'empresa_antiga': 'Gloob', 'empresa_nova': 'Globo Animação'},
            {'nome': 'vanessa galvão', 'empresa_antiga': 'Multishow', 'empresa_nova': 'GLOBO Humor'},
            {'nome': 'marina bouças', 'empresa_antiga': 'Play 9', 'empresa_nova': 'Play Action'},
            {'nome': 'amanda kadobayashi', 'empresa_antiga': 'PinguimTV', 'empresa_nova': 'Pinguim Content'},
        ],
        'participantes_novos': [      # [{nome, empresa, cargo, email, indisp_slots: [[dia,slot]]}]
            {
                'nome': 'Vinícius Neris', 'empresa': 'Globo Filmes',
                'cargo': 'Analista de conteúdo', 'email': 'vinicius.neris@g.globo',
                'obs': 'Disponível apenas em 29/05',
                'indisp_slots': [[d, s] for d in ['27/05 (Qua)', '28/05 (Qui)'] for s in SLOTS],
            },
            {
                'nome': 'Vinícius Lobo', 'empresa': 'Globo Filmes',
                'cargo': 'Curador de Conteúdo', 'email': 'vinicius.lobo@g.globo',
                'obs': 'Disponível apenas em 29/05',
                'indisp_slots': [[d, s] for d in ['27/05 (Qua)', '28/05 (Qui)'] for s in SLOTS],
            },
        ],
    }


def save_overrides(data):
    with open(OVERRIDES_PATH, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


# ============== CARREGAR PARTICIPANTES ==============
def is_korea_participant(country_str, company_str, name_str=''):
    """Detecta se o participante é da Coreia."""
    haystack = f"{country_str} {company_str} {name_str}".lower()
    return any(k in haystack for k in ['coreia', 'korea', 'corea'])


def load_participants(overrides):
    df_raw = pd.read_excel(EXCEL_PATH, sheet_name='agenda2026', header=None)
    participants = {}
    company_groups = defaultdict(list)

    for idx in range(2, len(df_raw)):
        row = df_raw.iloc[idx]
        name = str(row.iloc[0]).strip() if pd.notna(row.iloc[0]) else ''
        company = str(row.iloc[1]).strip() if pd.notna(row.iloc[1]) else ''
        cargo = str(row.iloc[3]).strip() if pd.notna(row.iloc[3]) else ''
        obs = str(row.iloc[4]).strip() if pd.notna(row.iloc[4]) else ''
        email = str(row.iloc[5]).strip() if pd.notna(row.iloc[5]) else ''
        if not name or name == 'nan' or name.startswith('🌎') or name.startswith('🇧🇷') or name == 'NOME' or name == 'NUMERO DE REUNIÕES':
            continue
        availability = {}
        for day, cols in DAY_COL_RANGES.items():
            for slot_idx, col in enumerate(cols):
                val = str(row.iloc[col]).strip().lower() if pd.notna(row.iloc[col]) else ''
                if val == 'nan': val = ''
                availability[(day, SLOTS[slot_idx])] = val
        pid = f"{name}|{company}"
        participants[pid] = {'name': name, 'company': company, 'cargo': cargo,
                            'obs': obs, 'email': email, 'availability': availability,
                            'scheduled': {}}
        company_groups[company].append(pid)

    # Apply correcoes_empresa
    for c in overrides.get('correcoes_empresa', []):
        old_pid = None
        for pid, p in list(participants.items()):
            if c['nome'].lower() in p['name'].lower() and p['company'] == c['empresa_antiga']:
                old_pid = pid
                break
        if old_pid:
            p = participants.pop(old_pid)
            if old_pid in company_groups[c['empresa_antiga']]:
                company_groups[c['empresa_antiga']].remove(old_pid)
            p['company'] = c['empresa_nova']
            new_pid = f"{p['name']}|{p['company']}"
            participants[new_pid] = p
            company_groups[c['empresa_nova']].append(new_pid)

    # Apply participantes_novos
    for p in overrides.get('participantes_novos', []):
        availability = {(d, s): '' for d in DAYS for s in SLOTS}
        for sk in p.get('indisp_slots', []):
            availability[(sk[0], sk[1])] = 'indisp'
        pid = f"{p['nome']}|{p['empresa']}"
        if pid not in participants:
            participants[pid] = {
                'name': p['nome'], 'company': p['empresa'], 'cargo': p.get('cargo', ''),
                'obs': p.get('obs', ''), 'email': p.get('email', ''),
                'availability': availability, 'scheduled': {},
            }
            company_groups[p['empresa']].append(pid)

    # Apply bloqueios_horario
    for b in overrides.get('bloqueios_horario', []):
        empresa_lower = b['empresa'].lower()
        target_pids = [pid for pid, p in participants.items() if empresa_lower in p['company'].lower() or p['company'].lower() in empresa_lower]
        if 'pessoa' in b and b['pessoa']:
            pessoa_lower = b['pessoa'].lower()
            target_pids = [pid for pid in target_pids if pessoa_lower in participants[pid]['name'].lower()]
        for pid in target_pids:
            participants[pid]['availability'][(b['dia'], b['slot'])] = 'indisp'

    return participants, company_groups


# ============== CARREGAR MATCHES ==============
def parse_matches_xlsx(xlsx_path):
    """Lê matches da planilha XLSX (formato novo). Lê TODAS as linhas com Empresa A e Empresa B."""
    df = pd.read_excel(xlsx_path, sheet_name=0, header=0)
    matches = []
    auto_num = 1
    for idx, row in df.iterrows():
        # Empresa A e B são os ÚNICOS campos obrigatórios
        company_a = str(row.get('Empresa A', '')).strip() if pd.notna(row.get('Empresa A')) else ''
        company_b = str(row.get('Empresa B', '')).strip() if pd.notna(row.get('Empresa B')) else ''
        if not company_a or company_a.lower() == 'nan': continue
        if not company_b or company_b.lower() == 'nan': continue

        # Número: usa coluna # se válida, senão auto-gera
        try:
            num = int(row['#']) if pd.notna(row.get('#')) else auto_num
        except (ValueError, KeyError, TypeError):
            num = auto_num
        auto_num = max(auto_num, num) + 1

        # Tier: "Tier 1 — Int×Nac" -> "Tier 1". Default "Tier 3" se ausente.
        tier_raw = str(row.get('Tier', '')).strip() if pd.notna(row.get('Tier')) else ''
        tm = re.match(r'(Tier \d)', tier_raw)
        tier = tm.group(1) if tm else 'Tier 3'

        # Score
        try:
            score = int(float(row.get('Score', 0))) if pd.notna(row.get('Score')) else 50
        except (ValueError, TypeError):
            score = 50

        contato_a = str(row.get('Contato A', '')).strip() if pd.notna(row.get('Contato A')) else ''
        contato_b = str(row.get('Contato B', '')).strip() if pd.notna(row.get('Contato B')) else ''

        def split_reps(s):
            if not s or s.lower() == 'nan': return []
            return [r.strip() for r in re.split(r'\s+e\s+|,\s*|/\s*|\s+&\s+|;\s*', s) if r.strip() and r.strip().lower() != 'nan']

        matches.append({
            'num': num, 'company_a': company_a, 'company_b': company_b,
            'reps_a': split_reps(contato_a), 'reps_b': split_reps(contato_b),
            'tier': tier, 'score': score,
        })
    return matches


def parse_matches_pdf():
    reader = PyPDF2.PdfReader(PDF_PATH)
    text = '\n'.join(p.extract_text() for p in reader.pages)
    blocks = re.split(r'\n(\d{1,3})\.\s+', text)
    matches = []
    for i in range(1, len(blocks)-1, 2):
        num = int(blocks[i])
        block = blocks[i+1]
        lines = block.strip().split('\n')
        first = lines[0] if lines else ''
        parts = re.split(r'×|x', first, maxsplit=1)
        if len(parts) < 2: continue
        company_a = parts[0].strip()
        company_b = ''
        score, tier = 0, 'Tier 3'
        reps_a, reps_b = [], []
        for line in lines:
            if 'Tier' in line and 'Score' in line:
                sm = re.search(r'Score\s+(\d+)', line)
                if sm: score = int(sm.group(1))
                tm = re.search(r'(Tier \d)', line)
                if tm: tier = tm.group(1)
        for line in lines:
            ls = line.strip()
            if ls.startswith('A:'):
                m = re.match(r'A:\s*(.+?)\s*\(', ls)
                if m: company_a = m.group(1).strip()
                rm = re.search(r'—\s*([^—]+)$', ls)
                if rm: reps_a = [r.strip() for r in re.split(r'\s+e\s+|,\s*|/\s*', rm.group(1))]
            elif ls.startswith('B:'):
                m = re.match(r'B:\s*(.+?)\s*\(', ls)
                if m: company_b = m.group(1).strip()
                rm = re.search(r'—\s*([^—]+)$', ls)
                if rm: reps_b = [r.strip() for r in re.split(r'\s+e\s+|,\s*|/\s*', rm.group(1))]
        matches.append({'num': num, 'company_a': company_a, 'company_b': company_b,
                       'reps_a': reps_a, 'reps_b': reps_b, 'tier': tier, 'score': score})
    return matches


def load_matches(overrides):
    # Prefere XLSX (novo formato). Se não tiver, cai pro PDF.
    xlsx_path = find_matches_xlsx()
    if xlsx_path:
        matches = parse_matches_xlsx(xlsx_path)
    else:
        matches = parse_matches_pdf()

    # Cancelar matches
    cancelled = overrides.get('matches_cancelados', [])
    def is_cancelled(m):
        for c in cancelled:
            a, b = c['empresa_a'].lower().strip(), c['empresa_b'].lower().strip()
            ma, mb = m['company_a'].lower().strip(), m['company_b'].lower().strip()
            if {a, b} == {ma, mb}:
                return True
        return False
    matches = [m for m in matches if not is_cancelled(m)]

    # Adicionar matches extras
    next_num = max((m['num'] for m in matches), default=0) + 1
    for e in overrides.get('matches_extras', []):
        matches.append({
            'num': next_num,
            'company_a': e['empresa_a'], 'company_b': e['empresa_b'],
            'reps_a': e.get('reps_a', []), 'reps_b': e.get('reps_b', []),
            'tier': e.get('tier', 'Tier 3'), 'score': e.get('score', 50),
        })
        next_num += 1

    return matches


# ============== RESOLVER MATCHES ==============
def resolve_matches(matches, participants, company_groups):
    SPECIAL = {'sidênia freire pereira': 'Sidênia Freire', 'vitor gabriel da silva': 'Vitor Gabriel',
               'amina sophia nogueira': 'Amina Nogueira', 'tiago campany': 'Tiago Camapny'}

    import unicodedata
    def strip_accents(s):
        return ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn').lower().strip()

    def find_pids(name, company_hint=''):
        nl = strip_accents(name)
        if not nl: return []
        # PASSO 1: busca todos os candidatos pelo nome (sem acento)
        all_candidates = []
        for pid, p in participants.items():
            pn = strip_accents(p['name'])
            if nl in pn or pn in nl:
                all_candidates.append(pid)
        # Se nenhum match completo, tenta primeiro+último nome
        if not all_candidates:
            parts = nl.split()
            if len(parts) >= 2:
                first, last = parts[0], parts[-1]
                for pid, p in participants.items():
                    pn = strip_accents(p['name'])
                    if first in pn and last in pn:
                        all_candidates.append(pid)
        # Se ainda nada, primeiro nome dentro do company_hint
        if not all_candidates and company_hint:
            ch = strip_accents(company_hint)
            parts = nl.split()
            if parts:
                first = parts[0]
                for pid, p in participants.items():
                    pc = strip_accents(p['company'])
                    if (ch in pc or pc in ch) and first in strip_accents(p['name']):
                        all_candidates.append(pid)
        if not all_candidates: return []
        # PASSO 2: se há múltiplos candidatos, filtra pelo company_hint (se houver)
        if len(all_candidates) > 1 and company_hint:
            ch = strip_accents(company_hint)
            # Tier 1: empresa exata
            tier1 = [pid for pid in all_candidates
                     if strip_accents(participants[pid]['company']) == ch]
            if tier1: return tier1
            # Tier 2: substring
            tier2 = [pid for pid in all_candidates
                     if ch in strip_accents(participants[pid]['company'])
                     or strip_accents(participants[pid]['company']) in ch]
            if tier2: return tier2
            # Tier 3: pelo menos uma palavra em comum (ex: "Kromaki 1" e "Kromaki - Fátima")
            ch_words = set(ch.split())
            tier3 = []
            for pid in all_candidates:
                pc_words = set(strip_accents(participants[pid]['company']).split())
                if ch_words & pc_words: tier3.append(pid)
            if tier3: return tier3
        return all_candidates

    # Companion map — APENAS quando a obs explicitamente fala "acompanha".
    # Antes: pareava automaticamente quando 2 pessoas tinham mesma disponibilidade,
    # mas isso causava bugs (ex: Kromaki Zico recebendo reuniões da Kromaki Fátima).
    COMPANION = {}
    for company, pids in company_groups.items():
        if len(pids) >= 2:
            for pid in pids:
                obs = participants[pid]['obs'].lower()
                if 'acompanha' in obs:
                    for o in pids:
                        if o != pid: COMPANION[pid] = o

    resolved = []
    unresolved_log = []
    for m in matches:
        pa, pb = set(), set()
        for r in m['reps_a']: pa.update(find_pids(r, m['company_a']))
        for r in m['reps_b']: pb.update(find_pids(r, m['company_b']))
        # Tenta SPECIAL antes de fallback genérico
        if not pa:
            for r in m['reps_a']:
                mp = SPECIAL.get(r.lower().strip())
                if mp: pa.update(find_pids(mp))
        if not pb:
            for r in m['reps_b']:
                mp = SPECIAL.get(r.lower().strip())
                if mp: pb.update(find_pids(mp))
        # Fallback por empresa SOMENTE se não achamos ninguém pelo nome.
        # Estratégia em camadas:
        # 1) Tenta match exato (case-insensitive)
        # 2) Tenta cada parte de nomes combinados (ex: "Canal Rural / BR IN TV")
        # 3) Não usa substring genérico (evita puxar Kromaki Zico em match da Kromaki Fátima)
        def fallback_by_company(target_company):
            target_lower = target_company.lower().strip()
            results = set()
            # Camada 1: match exato
            for pid, p in participants.items():
                if p['company'].lower().strip() == target_lower:
                    results.add(pid)
            if results: return results
            # Camada 2: divide em partes (separador / ou &)
            parts = [pt.strip().lower() for pt in re.split(r'\s*[/&]\s*', target_company) if pt.strip()]
            if len(parts) > 1:
                for part in parts:
                    for pid, p in participants.items():
                        if p['company'].lower().strip() == part:
                            results.add(pid)
            return results

        if not pa:
            pa.update(fallback_by_company(m['company_a']))
        if not pb:
            pb.update(fallback_by_company(m['company_b']))
        # Companion: só pra quem tem obs explícita "acompanha"
        for pid in list(pa):
            if pid in COMPANION and COMPANION[pid] not in pa: pa.add(COMPANION[pid])
        for pid in list(pb):
            if pid in COMPANION and COMPANION[pid] not in pb: pb.add(COMPANION[pid])
        if pa and pb:
            resolved.append({**m, 'pids_a': pa, 'pids_b': pb, 'all_pids': pa | pb})
        else:
            unresolved_log.append((m['num'], m['company_a'], m['company_b'],
                                   'pa vazio' if not pa else 'pb vazio'))
    if unresolved_log:
        print(f"[resolve_matches] {len(unresolved_log)} match(es) não resolvido(s):")
        for num, ca, cb, why in unresolved_log[:5]:
            print(f"  • #{num} {ca} x {cb} ({why})")
    return resolved


# ============== AGENDAR ==============
def schedule_matches(matches, participants):
    def is_avail(pid, day, slot):
        sk = (day, slot)
        return participants[pid]['availability'].get(sk, '') not in BLOCKED_STATUSES and sk not in participants[pid]['scheduled']

    def all_avail(pids, day, slot):
        return all(is_avail(pid, day, slot) for pid in pids)

    tier_pri = {'Tier 1': 0, 'Tier 2': 1, 'Tier 3': 2}
    matches.sort(key=lambda m: (tier_pri.get(m['tier'], 3), -m['score']))

    schedule = defaultdict(list)
    scheduled, unscheduled = [], []
    load = defaultdict(int)

    for m in matches:
        best, best_score = None, None
        for day in DAYS:
            for slot in SLOTS:
                if len(schedule[(day, slot)]) >= cap_for(day, slot): continue
                if not all_avail(m['all_pids'], day, slot): continue
                loads = [load[pid] for pid in m['all_pids']]
                slot_idx = SLOTS.index(slot)
                day_meet = sum(1 for sm in scheduled if sm['day'] == day)
                slot_use = len(schedule[(day, slot)])
                cons_pen = 0
                for pid in m['all_pids']:
                    if slot_idx > 0 and (day, SLOTS[slot_idx-1]) in participants[pid]['scheduled']:
                        cons_pen += 1
                    if slot_idx < len(SLOTS)-1 and (day, SLOTS[slot_idx+1]) in participants[pid]['scheduled']:
                        cons_pen += 1
                score = (day_meet, slot_use, cons_pen, max(loads), sum(loads))
                if best_score is None or score < best_score:
                    best_score = score; best = (day, slot)
        if best:
            d, s = best
            schedule[(d, s)].append(m)
            for pid in m['all_pids']:
                participants[pid]['scheduled'][(d, s)] = m['num']
                load[pid] += 1
            scheduled.append({**m, 'day': d, 'slot': s})
        else:
            unscheduled.append(m)

    # Swap rescue
    rescued = []
    still = []
    for m in unscheduled:
        ok = False
        for day in DAYS:
            if ok: break
            for slot in SLOTS:
                sk = (day, slot)
                if not all(participants[pid]['availability'].get(sk, '') not in BLOCKED_STATUSES for pid in m['all_pids']):
                    continue
                blocking = []
                for pid in m['all_pids']:
                    if sk in participants[pid]['scheduled']:
                        bn = participants[pid]['scheduled'][sk]
                        bm = next((sm for sm in scheduled if sm['num'] == bn), None)
                        if bm and bm not in blocking:
                            blocking.append(bm)
                if not blocking and len(schedule[sk]) >= cap_for(day, slot): continue
                # Garante que após mover blocking + adicionar m, slot não excede capacidade
                if (len(schedule[sk]) - len(blocking) + 1) > cap_for(day, slot): continue

                relocs = []
                fail = False
                # Reserva tentativa: pids -> set(slot) e contagem por slot
                tentative_pid_slots = defaultdict(set)
                tentative_slot_count = defaultdict(int)
                for bm in blocking:
                    ns = None
                    for d2 in DAYS:
                        for s2 in SLOTS:
                            sk2 = (d2, s2)
                            if sk2 == sk: continue
                            current = len(schedule[sk2]) + tentative_slot_count[sk2]
                            if current >= cap_for(d2, s2): continue
                            if not all(participants[pid]['availability'].get(sk2, '') not in BLOCKED_STATUSES for pid in bm['all_pids']): continue
                            if any(sk2 in participants[pid]['scheduled'] for pid in bm['all_pids']): continue
                            # Evita conflito com OUTRAS realocações nesta tentativa
                            if any(sk2 in tentative_pid_slots[pid] for pid in bm['all_pids']): continue
                            ns = sk2; break
                        if ns: break
                    if not ns: fail = True; break
                    relocs.append((bm, ns))
                    for pid in bm['all_pids']: tentative_pid_slots[pid].add(ns)
                    tentative_slot_count[ns] += 1
                if fail: continue
                for bm, (nd, nsl) in relocs:
                    ok_key = (bm['day'], bm['slot'])
                    schedule[ok_key] = [x for x in schedule[ok_key] if x['num'] != bm['num']]
                    for pid in bm['all_pids']:
                        if ok_key in participants[pid]['scheduled']:
                            del participants[pid]['scheduled'][ok_key]
                    schedule[(nd, nsl)].append(bm)
                    bm['day'], bm['slot'] = nd, nsl
                    for pid in bm['all_pids']: participants[pid]['scheduled'][(nd, nsl)] = bm['num']
                schedule[sk].append(m)
                for pid in m['all_pids']:
                    participants[pid]['scheduled'][sk] = m['num']; load[pid] += 1
                scheduled.append({**m, 'day': day, 'slot': slot})
                rescued.append(m); ok = True; break
        if not ok: still.append(m)

    # SANITY CHECK: garante consistência entre `scheduled` e `schedule[]`.
    # Se algo dessincronizou (bug em rescue ou edição posterior), remove órfãos.
    valid_scheduled = []
    seen_nums_in_schedule = set()
    for sk_key, mlist in schedule.items():
        for sm in mlist:
            seen_nums_in_schedule.add(sm['num'])
    for sm in scheduled:
        sk = (sm['day'], sm['slot'])
        in_correct_slot = any(x['num'] == sm['num'] for x in schedule.get(sk, []))
        if in_correct_slot:
            valid_scheduled.append(sm)
        else:
            # Match em scheduled mas não no slot informado. Tenta achar onde está.
            actual_slot = None
            for slot_key, mlist in schedule.items():
                if any(x['num'] == sm['num'] for x in mlist):
                    actual_slot = slot_key; break
            if actual_slot:
                # Atualiza day/slot pra refletir realidade do schedule[]
                sm['day'], sm['slot'] = actual_slot
                valid_scheduled.append(sm)
            else:
                # Match não está em lugar nenhum do schedule. Move pra unscheduled.
                still.append(sm)
    scheduled = valid_scheduled

    return scheduled, still, schedule


# ============== GERAR EXCELS ==============
TITLE_FONT = Font(name='Arial', size=16, bold=True, color='FFFFFF')
SUBTITLE_FONT = Font(name='Arial', size=11, italic=True, color='FFFFFF')
HEADER_FONT = Font(name='Arial', size=10, bold=True, color='FFFFFF')
DAY_FONT = Font(name='Arial', size=12, bold=True, color='FFFFFF')
CELL_FONT = Font(name='Arial', size=10)
BOLD_FONT = Font(name='Arial', size=10, bold=True)
SLOT_FONT = Font(name='Arial', size=10, bold=True, color='1F4E79')
PRIMARY_FILL = PatternFill('solid', fgColor='1F4E79')
DAY_FILLS = {'27/05 (Qua)': PatternFill('solid', fgColor='2E75B6'),
             '28/05 (Qui)': PatternFill('solid', fgColor='548235'),
             '29/05 (Sex)': PatternFill('solid', fgColor='C65911')}
ROW_FILLS = {'27/05 (Qua)': PatternFill('solid', fgColor='DEEBF7'),
             '28/05 (Qui)': PatternFill('solid', fgColor='E2EFDA'),
             '29/05 (Sex)': PatternFill('solid', fgColor='FCE4D6')}
ALT_FILL = PatternFill('solid', fgColor='F8F9FA')
BORDER = Border(left=Side(style='thin', color='D0D0D0'), right=Side(style='thin', color='D0D0D0'),
               top=Side(style='thin', color='D0D0D0'), bottom=Side(style='thin', color='D0D0D0'))


def safe_filename(s):
    s = re.sub(r'[<>:"/\\|?*]', '_', s)
    s = re.sub(r'\s+', '_', s.strip())
    return s[:80]


def gerar_agenda_geral(schedule, participants, output_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Agenda Sala Transforma"
    ws.merge_cells('A1:F1')
    ws['A1'] = 'SALA TRANSFORMA — RIO2C 2026'
    ws['A1'].font = TITLE_FONT; ws['A1'].fill = PRIMARY_FILL
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 32
    ws.merge_cells('A2:F2')
    ws['A2'] = 'Agenda de Matchmaking — 27, 28 e 29 de maio'
    ws['A2'].font = SUBTITLE_FONT; ws['A2'].fill = PRIMARY_FILL
    ws['A2'].alignment = Alignment(horizontal='center')
    ws.row_dimensions[2].height = 22

    row = 4
    for day in DAYS:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        c = ws.cell(row=row, column=1, value=f'  {day.upper()}')
        c.font = DAY_FONT; c.fill = DAY_FILLS[day]
        c.alignment = Alignment(horizontal='left', vertical='center')
        ws.row_dimensions[row].height = 26
        row += 1
        for col, h in enumerate(['Horário', 'Mesa', 'Empresa A', 'Representantes A', 'Empresa B', 'Representantes B'], 1):
            c = ws.cell(row=row, column=col, value=h)
            c.font = HEADER_FONT; c.fill = DAY_FILLS[day]; c.border = BORDER
            c.alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[row].height = 22
        row += 1
        for slot in SLOTS:
            meetings = sorted(schedule[(day, slot)], key=lambda x: x['company_a'].lower())
            if not meetings:
                c = ws.cell(row=row, column=1, value=slot)
                c.font = SLOT_FONT; c.fill = ROW_FILLS[day]
                c.alignment = Alignment(horizontal='center', vertical='center')
                ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
                c2 = ws.cell(row=row, column=2, value='— sem reuniões —')
                c2.font = Font(name='Arial', size=9, italic=True, color='999999'); c2.fill = ROW_FILLS[day]
                c2.alignment = Alignment(horizontal='center')
                for cc in range(1, 7): ws.cell(row=row, column=cc).border = BORDER
                row += 1
            else:
                start = row
                for idx, m in enumerate(meetings, 1):
                    reps_a = ', '.join(participants[pid]['name'] for pid in m['pids_a'])
                    reps_b = ', '.join(participants[pid]['name'] for pid in m['pids_b'])
                    data = [slot if idx == 1 else '', f'Mesa {idx}', m['company_a'], reps_a, m['company_b'], reps_b]
                    fill = ROW_FILLS[day] if idx % 2 == 1 else ALT_FILL
                    for col, val in enumerate(data, 1):
                        c = ws.cell(row=row, column=col, value=val)
                        c.font = SLOT_FONT if col == 1 else (BOLD_FONT if col == 2 else CELL_FONT)
                        c.fill = fill
                        c.alignment = Alignment(horizontal='center' if col in (1, 2) else 'left',
                                              vertical='center', wrap_text=True)
                        c.border = BORDER
                    row += 1
                if len(meetings) > 1:
                    ws.merge_cells(start_row=start, start_column=1, end_row=row-1, end_column=1)
        row += 1

    ws.column_dimensions['A'].width = 14; ws.column_dimensions['B'].width = 9
    for col in 'CDEF': ws.column_dimensions[col].width = 32
    ws.freeze_panes = 'A4'
    wb.save(output_path)


def gerar_agenda_empresa(company, scheduled, participants, company_groups, out_dir):
    pids = company_groups.get(company, [])
    if not pids: return None
    pids_set = set(pids)
    meetings = []
    for m in scheduled:
        # Verifica APENAS via pids_a/pids_b (que vêm de find_pids do nome do rep).
        # Companion automático foi desligado, então isso é confiável.
        in_a = pids_set & set(m['pids_a'])
        in_b = pids_set & set(m['pids_b'])
        if not in_a and not in_b: continue
        is_a = bool(in_a)
        our_pids_in_match = list(in_a if is_a else in_b)
        partner = m['company_b'] if is_a else m['company_a']
        partner_pids = m['pids_b'] if is_a else m['pids_a']
        meetings.append({'day': m['day'], 'slot': m['slot'],
                        'partner_company': partner,
                        'partner_reps': ', '.join(participants[pid]['name'] for pid in partner_pids),
                        'our_reps': ', '.join(participants[pid]['name'] for pid in our_pids_in_match)})
    if not meetings: return None
    meetings.sort(key=lambda x: (DAYS.index(x['day']), SLOTS.index(x['slot'])))
    wb = Workbook(); ws = wb.active; ws.title = "Minha Agenda"
    ws.merge_cells('A1:E1')
    ws['A1'] = f'AGENDA — {company.upper()}'
    ws['A1'].font = TITLE_FONT; ws['A1'].fill = PRIMARY_FILL
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 32
    ws.merge_cells('A2:E2')
    ws['A2'] = f'Sala Transforma — Rio2C 2026  •  {len(meetings)} reuniões agendadas'
    ws['A2'].font = SUBTITLE_FONT; ws['A2'].fill = PRIMARY_FILL
    ws['A2'].alignment = Alignment(horizontal='center')
    ws.row_dimensions[2].height = 22
    reps_text = ', '.join(sorted(set(participants[pid]['name'] for pid in pids)))
    ws.merge_cells('A3:E3')
    ws['A3'] = f'Representantes: {reps_text}'
    ws['A3'].font = Font(name='Arial', size=10, italic=True, color='555555')
    ws['A3'].alignment = Alignment(horizontal='center')
    ws.row_dimensions[3].height = 20
    row = 5; current_day = None
    for me in meetings:
        if me['day'] != current_day:
            current_day = me['day']
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
            c = ws.cell(row=row, column=1, value=f'  {current_day.upper()}')
            c.font = DAY_FONT; c.fill = DAY_FILLS[current_day]
            c.alignment = Alignment(horizontal='left', vertical='center')
            ws.row_dimensions[row].height = 26; row += 1
            for col, h in enumerate(['Horário', 'Empresa', 'País / Tipo', 'Representante(s)', 'Nossa Equipe'], 1):
                c = ws.cell(row=row, column=col, value=h)
                c.font = HEADER_FONT; c.fill = DAY_FILLS[current_day]; c.border = BORDER
                c.alignment = Alignment(horizontal='center', vertical='center')
            ws.row_dimensions[row].height = 22; row += 1
        partner_pid = next((pid for pid, p in participants.items() if p['company'].lower() == me['partner_company'].lower()), None)
        info = participants[partner_pid]['cargo'] if partner_pid else ''
        data = [me['slot'], me['partner_company'], info, me['partner_reps'], me['our_reps']]
        fill = ROW_FILLS[me['day']]
        for col, val in enumerate(data, 1):
            c = ws.cell(row=row, column=col, value=val)
            c.font = SLOT_FONT if col == 1 else (BOLD_FONT if col == 2 else CELL_FONT)
            c.fill = fill
            c.alignment = Alignment(horizontal='center' if col == 1 else 'left', vertical='center', wrap_text=True)
            c.border = BORDER
        ws.row_dimensions[row].height = 30; row += 1
    ws.column_dimensions['A'].width = 14
    for col, w in zip('BCDE', [32, 28, 32, 32]): ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A5'
    fname = safe_filename(company) + '.xlsx'
    out = os.path.join(out_dir, fname)
    wb.save(out)
    return out


# ============== AGENDA DETALHADA COM ESTATÍSTICAS ==============
def gerar_agenda_detalhada(scheduled_matches, unscheduled_matches, resolved_matches,
                            schedule, participants, output_path):
    """Gera o arquivo completo com 5 abas: Geral, Por Participante, Não Agendados, Estatísticas, Carga."""
    wb = Workbook()

    header_fill = PatternFill('solid', fgColor='1F4E79')
    header_font = Font(bold=True, color='FFFFFF', name='Arial', size=10)
    day_fills = {'27/05 (Qua)': PatternFill('solid', fgColor='D6E4F0'),
                 '28/05 (Qui)': PatternFill('solid', fgColor='E2EFDA'),
                 '29/05 (Sex)': PatternFill('solid', fgColor='FCE4D6')}
    slot_fills = {'27/05 (Qua)': PatternFill('solid', fgColor='EBF1F8'),
                  '28/05 (Qui)': PatternFill('solid', fgColor='F0F7EC'),
                  '29/05 (Sex)': PatternFill('solid', fgColor='FEF2EB')}
    border = Border(left=Side(style='thin', color='BFBFBF'), right=Side(style='thin', color='BFBFBF'),
                   top=Side(style='thin', color='BFBFBF'), bottom=Side(style='thin', color='BFBFBF'))
    cell_font = Font(name='Arial', size=9)
    bold_font = Font(name='Arial', size=9, bold=True)
    title_font = Font(name='Arial', size=14, bold=True, color='1F4E79')

    # ===== Aba 1: Agenda Geral =====
    ws1 = wb.active; ws1.title = "Agenda Geral"
    ws1.merge_cells('A1:H1')
    ws1['A1'] = 'Sala Transforma — Rio2C 2026 — Agenda de Matchmaking'
    ws1['A1'].font = title_font
    ws1['A2'] = f'{len(scheduled_matches)} reuniões agendadas de {len(resolved_matches)} matches aprovados'
    ws1['A2'].font = Font(name='Arial', size=10, italic=True, color='666666')
    row = 4
    for day in DAYS:
        ws1.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        cell = ws1.cell(row=row, column=1, value=f'Sala Transforma — {day}')
        cell.font = Font(name='Arial', size=12, bold=True, color='FFFFFF')
        cell.fill = header_fill; cell.alignment = Alignment(horizontal='center'); row += 1
        for col, h in enumerate(['Slot', 'Mesa', 'Match #', 'Empresa A', 'Empresa B', 'Representantes A', 'Representantes B', 'Score'], 1):
            cell = ws1.cell(row=row, column=col, value=h)
            cell.font = header_font; cell.fill = day_fills[day]
            cell.alignment = Alignment(horizontal='center', wrap_text=True); cell.border = border
        row += 1
        for slot in SLOTS:
            meetings = schedule[(day, slot)]
            if not meetings:
                cell = ws1.cell(row=row, column=1, value=slot)
                cell.font = bold_font; cell.fill = slot_fills[day]; cell.border = border
                cell = ws1.cell(row=row, column=2, value='—')
                cell.font = cell_font; cell.fill = slot_fills[day]
                cell.alignment = Alignment(horizontal='center'); cell.border = border
                for c in range(3, 9):
                    cell = ws1.cell(row=row, column=c, value='')
                    cell.fill = slot_fills[day]; cell.border = border
                row += 1
            else:
                for mesa_idx, m in enumerate(meetings, 1):
                    reps_a = [participants[pid]['name'] for pid in m['pids_a']]
                    reps_b = [participants[pid]['name'] for pid in m['pids_b']]
                    data = [slot if mesa_idx == 1 else '', f'Mesa {mesa_idx}', m['num'],
                           m['company_a'], m['company_b'], ', '.join(reps_a), ', '.join(reps_b), m['score']]
                    for col, val in enumerate(data, 1):
                        cell = ws1.cell(row=row, column=col, value=val)
                        cell.font = bold_font if col == 1 else cell_font
                        cell.fill = slot_fills[day]; cell.border = border
                        cell.alignment = Alignment(horizontal='center') if col in (1, 2, 3, 8) else Alignment(wrap_text=True)
                    row += 1
        row += 1
    for col, w in zip('ABCDEFGH', [14, 10, 10, 30, 30, 35, 35, 8]):
        ws1.column_dimensions[col].width = w

    # ===== Aba 2: Agenda por Participante =====
    ws2 = wb.create_sheet("Agenda por Participante")
    ws2.merge_cells('A1:I1')
    ws2['A1'] = 'Agenda Individual — Sala Transforma Rio2C 2026'
    ws2['A1'].font = title_font
    headers = ['Participante', 'Empresa', 'Cargo']
    for day in DAYS:
        for slot in SLOTS:
            headers.append(f'{day}\n{slot}')
    row = 3
    for col, h in enumerate(headers, 1):
        cell = ws2.cell(row=row, column=col, value=h)
        cell.font = header_font; cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', wrap_text=True, vertical='center'); cell.border = border
    ws2.row_dimensions[row].height = 40
    row = 4
    sorted_pids = sorted(participants.keys(), key=lambda pid: (participants[pid]['company'], participants[pid]['name']))
    meeting_fill = PatternFill('solid', fgColor='C6EFCE')
    blocked_fill = PatternFill('solid', fgColor='FFC7CE')
    buffer_fill = PatternFill('solid', fgColor='FFEB9C')
    free_fill = PatternFill('solid', fgColor='FFFFFF')
    for pid in sorted_pids:
        p = participants[pid]
        has_meeting = bool(p['scheduled'])
        has_avail = any(p['availability'].get((d, s), '') not in BLOCKED_STATUSES for d in DAYS for s in SLOTS)
        if not has_avail and not has_meeting: continue
        ws2.cell(row=row, column=1, value=p['name']).font = cell_font
        ws2.cell(row=row, column=1).border = border
        ws2.cell(row=row, column=2, value=p['company']).font = cell_font
        ws2.cell(row=row, column=2).border = border
        ws2.cell(row=row, column=3, value=p['cargo']).font = cell_font
        ws2.cell(row=row, column=3).border = border
        col = 4
        for day in DAYS:
            for slot in SLOTS:
                sk = (day, slot)
                cell = ws2.cell(row=row, column=col)
                cell.font = Font(name='Arial', size=7)
                cell.alignment = Alignment(horizontal='center', wrap_text=True); cell.border = border
                if sk in p['scheduled']:
                    mn = p['scheduled'][sk]
                    mi = next((m for m in scheduled_matches if m['num'] == mn), None)
                    if mi:
                        oc = mi['company_b'] if pid in mi['pids_a'] else mi['company_a']
                        cell.value = f"#{mn}\n{oc}"
                    else:
                        cell.value = f"#{mn}"
                    cell.fill = meeting_fill
                else:
                    s = p['availability'].get(sk, '')
                    if s in ('indisp', 'viagem', 'painel', 'pitching'):
                        cell.value = s.upper(); cell.fill = blocked_fill
                    elif s == 'rodada':
                        cell.value = 'RODADA'; cell.fill = PatternFill('solid', fgColor='BDD7EE')
                    elif s == 'buffer':
                        cell.value = 'buffer'; cell.fill = buffer_fill
                    else:
                        cell.value = ''; cell.fill = free_fill
                col += 1
        row += 1
    ws2.column_dimensions['A'].width = 25
    ws2.column_dimensions['B'].width = 28
    ws2.column_dimensions['C'].width = 30
    for c in range(4, 4 + len(DAYS) * len(SLOTS)):
        ws2.column_dimensions[get_column_letter(c)].width = 16

    # ===== Aba 3: Não Agendados (se houver) =====
    if unscheduled_matches:
        ws3 = wb.create_sheet("Não Agendados")
        ws3.merge_cells('A1:G1')
        ws3['A1'] = 'Matches Aprovados Não Agendados'
        ws3['A1'].font = title_font
        for col, h in enumerate(['Match #', 'Tier', 'Score', 'Empresa A', 'Empresa B', 'Motivo Provável', 'Reps A / Reps B'], 1):
            cell = ws3.cell(row=3, column=col, value=h)
            cell.font = header_font; cell.fill = PatternFill('solid', fgColor='C00000')
            cell.alignment = Alignment(horizontal='center'); cell.border = border
        for idx, m in enumerate(unscheduled_matches):
            r = 4 + idx
            ws3.cell(row=r, column=1, value=m['num']).font = cell_font
            ws3.cell(row=r, column=1).border = border
            ws3.cell(row=r, column=1).alignment = Alignment(horizontal='center')
            ws3.cell(row=r, column=2, value=m['tier']).font = cell_font
            ws3.cell(row=r, column=2).border = border
            ws3.cell(row=r, column=3, value=m['score']).font = cell_font
            ws3.cell(row=r, column=3).border = border
            ws3.cell(row=r, column=3).alignment = Alignment(horizontal='center')
            ws3.cell(row=r, column=4, value=m['company_a']).font = cell_font
            ws3.cell(row=r, column=4).border = border
            ws3.cell(row=r, column=5, value=m['company_b']).font = cell_font
            ws3.cell(row=r, column=5).border = border
            ws3.cell(row=r, column=6, value="Sem slot livre em comum").font = cell_font
            ws3.cell(row=r, column=6).border = border
            ws3.cell(row=r, column=7, value=f"A: {', '.join(m['reps_a'])} | B: {', '.join(m['reps_b'])}").font = cell_font
            ws3.cell(row=r, column=7).border = border
        for col, w in zip('ABCDEFG', [10, 18, 8, 30, 30, 40, 50]):
            ws3.column_dimensions[col].width = w

    # ===== Aba 4: Estatísticas =====
    ws4 = wb.create_sheet("Estatísticas")
    ws4['A1'] = 'Estatísticas do Agendamento'
    ws4['A1'].font = title_font
    stats = [('', ''),
             ('Total de Matches Aprovados', len(resolved_matches)),
             ('Matches Agendados', len(scheduled_matches)),
             ('Matches Não Agendados', len(unscheduled_matches)),
             ('Taxa de Agendamento', f'{len(scheduled_matches)/len(resolved_matches)*100:.1f}%' if resolved_matches else '0%'),
             ('', ''),
             ('Capacidade Total (3 dias)', sum(cap_for(d, s) for d in DAYS for s in SLOTS)),
             ('Ocupação Total', f'{len(scheduled_matches)}/{sum(cap_for(d, s) for d in DAYS for s in SLOTS)}'),
             ('', '')]
    for day in DAYS:
        stats.append((f'Reuniões {day}', sum(1 for m in scheduled_matches if m['day'] == day)))
    stats.append(('', '')); stats.append(('POR TIER', ''))
    tc = defaultdict(int)
    for m in scheduled_matches: tc[m['tier']] += 1
    for tier in ['Tier 1', 'Tier 2', 'Tier 3']:
        tot = sum(1 for m in resolved_matches if m['tier'] == tier)
        stats.append((f'{tier}: Agendados/Total', f'{tc.get(tier, 0)}/{tot}'))
    stats.append(('', '')); stats.append(('POR SLOT', ''))
    for day in DAYS:
        for slot in SLOTS:
            stats.append((f'{day} {slot}', f'{len(schedule[(day, slot)])}/{cap_for(day, slot)}'))
    row = 3
    for label, value in stats:
        ws4.cell(row=row, column=1, value=label).font = bold_font if label else cell_font
        ws4.cell(row=row, column=1).border = border
        ws4.cell(row=row, column=2, value=value).font = cell_font
        ws4.cell(row=row, column=2).border = border
        ws4.cell(row=row, column=2).alignment = Alignment(horizontal='center')
        row += 1
    ws4.column_dimensions['A'].width = 35; ws4.column_dimensions['B'].width = 20

    # ===== Aba 5: Carga por Participante =====
    ws5 = wb.create_sheet("Carga por Participante")
    ws5['A1'] = 'Reuniões Agendadas por Participante'
    ws5['A1'].font = title_font
    for col, h in enumerate(['Participante', 'Empresa', 'Reuniões Agendadas', 'Slots Livres Restantes', 'Slots Ocupados (Rodada/Outro)'], 1):
        cell = ws5.cell(row=3, column=col, value=h)
        cell.font = header_font; cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center'); cell.border = border
    row = 4
    for pid in sorted_pids:
        p = participants[pid]
        sc = len(p['scheduled']); fc = 0; bc = 0
        for d in DAYS:
            for s in SLOTS:
                sk = (d, s)
                if sk in p['scheduled']: continue
                st = p['availability'].get(sk, '')
                if st in BLOCKED_STATUSES: bc += 1
                else: fc += 1
        if sc == 0 and fc == 0 and bc == 0: continue
        ws5.cell(row=row, column=1, value=p['name']).font = cell_font
        ws5.cell(row=row, column=1).border = border
        ws5.cell(row=row, column=2, value=p['company']).font = cell_font
        ws5.cell(row=row, column=2).border = border
        for c, v in [(3, sc), (4, fc), (5, bc)]:
            ws5.cell(row=row, column=c, value=v).font = cell_font
            ws5.cell(row=row, column=c).alignment = Alignment(horizontal='center')
            ws5.cell(row=row, column=c).border = border
        row += 1
    for col, w in zip('ABCDE', [30, 30, 20, 22, 28]):
        ws5.column_dimensions[col].width = w

    wb.save(output_path)


# ============== EXECUÇÃO COMPLETA ==============
def run_full(out_general='Agenda_Equipe_SalaTransforma_2026.xlsx', out_dir='Agendas_Por_Empresa',
             out_detalhado='Agenda_SalaTransforma_Rio2C_2026.xlsx'):
    overrides = load_overrides()
    participants, company_groups = load_participants(overrides)
    matches = load_matches(overrides)

    # ====== BLOQUEAR COREIA NO DIA 27 ======
    # Detecta companies coreanas via planilha de matches (coluna País A/B) e bloqueia dia 27
    korean_companies = set()
    xlsx_path = find_matches_xlsx()
    if xlsx_path:
        try:
            df_m = pd.read_excel(xlsx_path, sheet_name=0, header=0)
            for _, row in df_m.iterrows():
                if is_korea_participant(str(row.get('País A', '')), str(row.get('Empresa A', ''))):
                    korean_companies.add(str(row.get('Empresa A', '')).strip())
                if is_korea_participant(str(row.get('País B', '')), str(row.get('Empresa B', ''))):
                    korean_companies.add(str(row.get('Empresa B', '')).strip())
        except Exception:
            pass
    # Bloqueia todos os slots do dia 27/05 para participantes dessas empresas
    for pid, p in participants.items():
        if any(kc and kc.lower() == p['company'].lower() for kc in korean_companies):
            for slot in SLOTS:
                p['availability'][('27/05 (Qua)', slot)] = 'indisp'

    resolved = resolve_matches(matches, participants, company_groups)
    scheduled, unscheduled, schedule = schedule_matches(resolved, participants)

    out_general_path = os.path.join(BASE_DIR, out_general)
    out_dir_path = os.path.join(BASE_DIR, out_dir)
    out_detalhado_path = os.path.join(BASE_DIR, out_detalhado)
    os.makedirs(out_dir_path, exist_ok=True)
    # Limpa arquivos antigos pra evitar cache de rodadas anteriores
    for old_file in os.listdir(out_dir_path):
        if old_file.endswith('.xlsx'):
            try:
                os.remove(os.path.join(out_dir_path, old_file))
            except Exception:
                pass
    gerar_agenda_geral(schedule, participants, out_general_path)
    gerar_agenda_detalhada(scheduled, unscheduled, resolved, schedule, participants, out_detalhado_path)
    # Empresas que realmente têm participantes em reuniões agendadas
    companies_with = set()
    for m in scheduled:
        for pid in (m['pids_a'] | m['pids_b']):
            companies_with.add(participants[pid]['company'])
    n_files = 0
    for c in sorted(companies_with):
        if gerar_agenda_empresa(c, scheduled, participants, company_groups, out_dir_path):
            n_files += 1

    return {
        'total_matches': len(resolved),
        'scheduled': len(scheduled),
        'unscheduled': unscheduled,
        'companies_with_meetings': len(companies_with),
        'files_generated': n_files,
        'output_general': out_general_path,
        'output_detalhado': out_detalhado_path,
        'output_dir': out_dir_path,
        'schedule': schedule,
        'scheduled_matches': scheduled,
        'participants': participants,
        'company_groups': company_groups,
    }


def get_all_companies():
    """Retorna lista de empresas únicas para uso nos formulários."""
    overrides = load_overrides()
    participants, company_groups = load_participants(overrides)
    return sorted(company_groups.keys())


def get_all_people():
    """Retorna lista de [nome, empresa] para dropdowns."""
    overrides = load_overrides()
    participants, _ = load_participants(overrides)
    return sorted([(p['name'], p['company']) for p in participants.values()])


if __name__ == '__main__':
    result = run_full()
    print(f"Agendados: {result['scheduled']}/{result['total_matches']}")
    print(f"Arquivos gerados: {result['files_generated']}")
    if result['unscheduled']:
        print(f"Não agendados: {len(result['unscheduled'])}")
        for m in result['unscheduled']:
            print(f"  • {m['company_a']} x {m['company_b']}")
